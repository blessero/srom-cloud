#!/usr/bin/env python3
"""volume_lists.py — lists of a whole volume, as a DOCX in the template's styles, placed into the volume's InDesign
document by hand like an article (same Word import preset).

  volume_lists.py noty <master.csv> --authors <autorzy.tsv> -o <out.docx>

noty — "Noty o autorach": one paragraph per person who authors a text in the parts of the volume before the last two
(MB 02.10.2026: parts I–II, or I–III when a theme title counts as part I; column `section_nr` of the master CSV; no
section_nr at all: every text) (master CSV `authors_struct`,
else `authors_display`), sorted by the register's key in Polish alphabetical order. The name as printed (degree
included) in the character style for bold, then " – " and the note (the note's own opening name and separator are
dropped); "Kontakt: …" after a line break; *…* in a note = italic. Texts as the register has them.
Reported, never guessed (Kanon § 0): a person not in the register or without a note (`[BRAK NOTY]` printed), a joint
note that opens with the name without a dash (printed as written, the opening in bold), and the register's own doubts
(column `uwagi`), and a capital after the dash (vol. 18 prints lower case; a proper name keeps it).
Last line: NOTY OK n  or  NOTY CHECK n (k to look at).
The contents page (spis treści) is not generated: it needs the part titles and the back-matter pages, which the master
CSV does not hold.
"""
import argparse, csv, json, os, re, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build import make_reference_docx  # noqa: E402  (the same styles as every article's DOCX)

ALPHA = "aąbcćdeęfghijklłmnńoóprsśtuvwxyzźż"
DEGREE = r"(?:prof\.\s+dr\s+hab\.|prof\.|dr\s+hab\.|dr)\s+"


def pl_key(s):
    return [ALPHA.index(c) if c in ALPHA else 100 + ord(c) for c in s.lower()]


def bare(name):
    """'Prof. dr hab. Anna G. Piotrowska (…)' -> 'Anna G. Piotrowska'"""
    return re.sub(r"\s*\(.*?\)\s*$", "", re.sub(r"^" + DEGREE, "", name.strip(), flags=re.I)).strip()


def persons(row):
    st = (row.get("authors_struct") or "").strip()
    if st:
        out = []
        for seg in st.split(";;"):
            f = [x.strip() for x in seg.split("|")]
            if len(f) >= 2 and (f[0] or f[1]):
                out.append(f"{f[0]} {f[1]}".strip())
        return out
    return [x.strip() for x in re.split(r",| i | and ", row.get("authors_display") or "") if x.strip()]


def split_note(person, note):
    """('clean', rest) when the note opens with '[degree] Name' and ',', '–' or '-' (dropped);
    ('named', end) when it opens with the name but goes on without a separator (a joint note); ('other', None)."""
    m = re.match(r"^(?:" + DEGREE + r")?" + re.escape(bare(person)), note, re.I)
    if not m:
        return "other", None
    s = re.match(r"\s*[,–-]\s*", note[m.end():])
    return ("clean", note[m.end() + s.end():]) if s else ("named", m.end())


def add_text(p, text, italic_style):
    for i, part in enumerate(re.split(r"\*([^*]+)\*", text)):
        if part:
            r = p.add_run(part)
            if i % 2:
                r.style = italic_style


def notes_parts(rows):
    """Kanon-free house rule (MB 02.10.2026, GEN-13): notes for the contributors of every part (`section_nr`, in order of
    appearance) except the last two (Recenzje, Opinie and the one after). A master CSV without section_nr: all rows."""
    parts = []
    for r in rows:
        k = (r.get("section_nr") or "").strip()
        if k and k not in parts:
            parts.append(k)
    if not parts:
        return rows
    keep = set(parts[:-2])
    return [r for r in rows if (r.get("section_nr") or "").strip() in keep]


def noty(a):
    import docx
    cfg = json.load(open(os.path.join(os.path.dirname(HERE), "config", "styles.json"), encoding="utf-8"))
    P, C = cfg["paragraph"], cfg["character"]
    reg = {bare(r["osoba"]).lower(): r for r in csv.DictReader(open(a.authors, encoding="utf-8"), delimiter="\t")}
    rows = list(csv.DictReader(open(a.csv, encoding="utf-8-sig")))
    rows = notes_parts(rows)
    seen, people = set(), []
    for row in rows:
        for n in persons(row):
            if n.lower() not in seen:
                seen.add(n.lower())
                people.append((n, reg.get(n.lower())))
    people.sort(key=lambda x: (pl_key(x[1]["sortuj"]) if x[1] else pl_key(x[0].split()[-1]), pl_key(x[0])))
    ref = os.path.join(tempfile.mkdtemp(), "reference.docx")
    make_reference_docx(cfg, ref)
    d = docx.Document(ref)
    for p in list(d.paragraphs):
        p._element.getparent().remove(p._element)
    report = []
    for name, r in people:
        p = d.add_paragraph(style=P["author_note"])
        shown = r["osoba"] if r else name
        p.add_run(re.sub(r"\s*\(.*?\)\s*$", "", shown)).style = C["speaker"]
        note = (r or {}).get("nota", "").strip()
        if not note:
            p.add_run(" – [BRAK NOTY]")
            report.append(f"{name}: " + ("not in the register" if not r else "no note in the register"))
            continue
        kind, val = split_note(shown, note)
        if kind == "named":                     # "Dr X i dr Y są autorami…": the note's own opening in bold
            p.runs[0].text = note[:val]
            p.add_run(" ")
            add_text(p, note[val:].lstrip(), C["italic"])
        else:
            p.add_run(" – ")
            add_text(p, val if kind == "clean" else note, C["italic"])
        if kind == "named":
            report.append(f"{name}: joint note (opens with the name, no dash) — printed as written")
        body = val if kind == "clean" else note
        if kind != "named" and re.match(r"[A-ZĄĆĘŁŃÓŚŹŻ][a-ząćęłńóśźż]", body):
            report.append(f"{name}: capital after the dash ({body.split()[0]}) — lower case unless a name")
        if r.get("uwagi"):
            report.append(f"{name}: register flags: {r['uwagi']}")
        if r.get("kontakt"):
            p.add_run().add_break()
            p.add_run("Kontakt: " + r["kontakt"])
    d.save(a.o)
    for line in report:
        print("  " + line)
    print(f"NOTY {'OK' if not report else 'CHECK'} {len(people)}" + (f" ({len(report)} to look at)" if report else "")
          + f" — {a.o}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    n = sub.add_parser("noty", help="Noty o autorach")
    n.add_argument("csv")
    n.add_argument("--authors", required=True)
    n.add_argument("-o", required=True)
    a = ap.parse_args()
    noty(a)


if __name__ == "__main__":
    main()
