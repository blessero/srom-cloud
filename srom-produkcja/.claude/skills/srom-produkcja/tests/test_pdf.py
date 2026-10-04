"""G7: PDF extractor on a born-digital article PDF made the way authors make them
(Word-style DOCX with real footnotes -> LibreOffice PDF: 11 pt justified body, 9 pt notes,
running head, page numbers, a footnote running over to the next page), plus a hand-set
page with line-end hyphenation. Ground truth = the markdown the PDF was generated from."""
import os, re, sys, shutil, subprocess, tempfile, difflib
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FX = os.path.join(ROOT, "tests", "fixtures", "pdf")
EX = os.path.join(ROOT, "scripts", "pdf_extract.py")
sys.path.insert(0, os.path.join(ROOT, "scripts"))
res = []

def t(name, ok, detail=""):
    res.append(ok); print(("PASS " if ok else "FAIL ") + name + ("" if ok else "\n   " + str(detail)[:1500]))

SOFFICE = shutil.which("soffice")
if not SOFFICE:
    print("SKIP LibreOffice part (Word-made PDF) — LibreOffice not available; hand-set pages below still run")
d = tempfile.mkdtemp()

if SOFFICE:
    # ---- build the PDF
    import docx
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ref = os.path.join(d, "ref.docx")
    open(ref, "wb").write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True).stdout)
    D = docx.Document(ref)
    for st in D.styles:
        n = getattr(st, "name", None)
        if n in ("Normal", "Body Text", "First Paragraph", "Compact"):
            st.font.size = Pt(11); st.font.name = "Liberation Serif"
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
        if n == "Body Text": st.paragraph_format.first_line_indent = Pt(18)
        if n == "Block Text":
            st.font.size = Pt(10); st.paragraph_format.left_indent = Pt(36); st.paragraph_format.right_indent = Pt(18)
            st.paragraph_format.space_before = Pt(6); st.paragraph_format.space_after = Pt(6)
        if n == "Footnote Text":
            st.font.size = Pt(9); st.font.name = "Liberation Serif"; st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    sec = D.sections[0]
    sec.header.paragraphs[0].text = "Journal of Test Studies 12 (2020)"
    r = sec.footer.paragraphs[0].add_run()
    for kind in ("begin", "instr", "separate", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = " PAGE "
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
    # footnotes numbered through the document, as in a Word-made article: pandoc's default reference.docx (3.8.3)
    # sets <w:numRestart w:val="eachSect"/>, which LibreOffice >= 26 turns into numbering restarting at every Heading 1
    for nr in sec._sectPr.xpath("w:footnotePr/w:numRestart"):
        nr.set(qn("w:val"), "continuous")
    D.save(ref)
    src = os.path.join(FX, "src.md")
    subprocess.run(["pandoc", src, "--reference-doc", ref, "-o", os.path.join(d, "art.docx")], check=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", d, os.path.join(d, "art.docx")], capture_output=True)
    pdf = os.path.join(d, "art.pdf")
    t("test PDF generated", os.path.exists(pdf))

    # ---- extract
    out = os.path.join(d, "ext.md")
    r = subprocess.run([sys.executable, EX, pdf, "-o", out], capture_output=True, text=True)
    rep = open(os.path.join(d, "ext_extract.md"), encoding="utf-8").read()
    t("extractor reports EXTRACT OK", r.returncode == 0 and "EXTRACT OK" in r.stdout, r.stdout + rep)
    ext = open(out, encoding="utf-8").read()
    srcmd = open(src, encoding="utf-8").read()

    def notes_of(md):
        return {int(m.group(1)): m.group(2).strip() for m in re.finditer(r"^\[\^(\d+)\]:\s*(.+)$", md, re.M)}
    def norm(s):
        return re.sub(r"\s+", " ", s.replace("\u00a0", " ")).strip()
    sn, en = notes_of(srcmd), notes_of(ext)
    t("note sequence 1..N contiguous, same count as source", sorted(en) == list(range(1, len(sn) + 1)), (sorted(en), len(sn)))
    seq = [int(x) for x in re.findall(r"\[\^(\d+)\](?!:)", ext)]
    t("markers appear once each, in order, == notes", seq == sorted(sn), seq)
    bad = [n for n in sn if norm(sn[n]) != norm(en.get(n, ""))]
    t("every note text identical to source incl. *italics* (after whitespace normalisation)", not bad,
      "\n".join(f"{n}: SRC {norm(sn[n])[:120]}\n   EXT {norm(en.get(n, ''))[:120]}" for n in bad[:3]))
    t("note running over a page boundary reassembled", "continues on page" in rep, rep)

    def body_paras(md):
        return [norm(re.sub(r"\[\^\d+\]", "", p)) for p in re.split(r"\n\s*\n", md) if p.strip() and not p.lstrip().startswith("[^")]
    sb, eb = body_paras(srcmd), body_paras(ext)
    t("same paragraph/heading/quote sequence", len(sb) == len(eb), (len(sb), len(eb)))
    diffs = [(a, b) for a, b in zip(sb, eb) if a != b]
    t("every body paragraph identical incl. italics, headings (#/##) and block quote (>)", not diffs,
      "\n".join(f"SRC {a[:140]}\nEXT {b[:140]}" for a, b in diffs[:3]))
    t("running head and page numbers dropped", "Journal of Test Studies" not in ext and not re.search(r"^\d+$", ext, re.M))
    t("marker positions identical to source", re.findall(r"(\S{0,12})\[\^(\d+)\](?!:)", ext) == re.findall(r"(\S{0,12})\[\^(\d+)\](?!:)", srcmd))

# ---- hand-set page: line-end hyphenation, italics across lines, marker glued to punctuation
import pymupdf
hp = os.path.join(d, "hyph.pdf")
# PyMuPDF's own Times (no font files: same on the Mac and in the sandbox; system TTFs can come back with
# nbsp for space and soft hyphen for hyphen, which is the font's cmap, not the extractor)
FONTS = {"rg": pymupdf.Font("tiro").buffer, "it": pymupdf.Font("tiit").buffer, "bd": pymupdf.Font("tibo").buffer}
doc = pymupdf.open(); pg = doc.new_page()
pg.insert_font(fontname="rg", fontbuffer=FONTS["rg"]); pg.insert_font(fontname="it", fontbuffer=FONTS["it"])
FR, FI = pymupdf.Font("tiro"), pymupdf.Font("tiit")
y = 100
lines = [("The history of the Ro-", "rg"), ("mani people is a self-", "rg"), ("evident topic of 1939–", "rg"), ("1945 research in the", "rg")]
for txt, fn in lines:
    pg.insert_text((72, y), txt, fontname=fn, fontsize=11); y += 14
pg.insert_text((72, y), "journal ", fontname="rg", fontsize=11)
x = 72 + FR.text_length("journal ", fontsize=11)
pg.insert_text((x, y), "Studia Romo-", fontname="it", fontsize=11); y += 14
pg.insert_text((72, y), "logica", fontname="it", fontsize=11)
x = 72 + FI.text_length("logica", fontsize=11)
pg.insert_text((x, y), " ends here.", fontname="rg", fontsize=11)
x += FR.text_length(" ends here.", fontsize=11)
pg.insert_text((x + 0.5, y - 4), "1", fontname="rg", fontsize=7)
pg.insert_text((72, 760), "1", fontname="rg", fontsize=6)
pg.insert_text((77, 764), "A note with a *star*.", fontname="rg", fontsize=9)
doc.save(hp)
r = subprocess.run([sys.executable, EX, hp, "-o", os.path.join(d, "h.md")], capture_output=True, text=True)
h = open(os.path.join(d, "h.md"), encoding="utf-8").read()
t("line-end hyphen removed: Ro-|mani -> Romani", "Romani people" in h, h)
t("prefix hyphen kept: self-|evident -> self-evident", "self-evident" in h, h)
t("range dash at line end kept, joined without space: 1939–|1945 -> 1939–1945", "1939–1945" in h, h)
t("italics continue across hyphenated line break: *Studia Romologica*", "*Studia Romologica*" in h, h)
t("marker after period on hand-set page", "here.[^1]" in h and "[^1]: A note" in h, h)

# ---- hand-set page 2: endnotes section, same-size indented quotation, hanging-indent reference list
doc = pymupdf.open(); pg = doc.new_page()
pg.insert_font(fontname="rg", fontbuffer=FONTS["rg"]); pg.insert_font(fontname="bd", fontbuffer=FONTS["bd"])
Y = [90]
def ln(x, txt, fn="rg", fs=11, gap=14):
    pg.insert_text((x, Y[0]), txt, fontname=fn, fontsize=fs); Y[0] += gap
def mark(after_x, txt, n):
    pg.insert_text((after_x + FR.text_length(txt, fontsize=11) + 0.5, Y[0] - 18), str(n), fontname="rg", fontsize=7)
ln(90, "A first paragraph opens with an indent and runs on for the full measure of")
ln(72, "the line and then ends somewhere short of the margin.")
mark(72, "the line and then ends somewhere short of the margin.", 1)
ln(108, "An indented quotation set at the body size runs")
ln(108, "across two lines between wider margins and ends.")
ln(90, "The next paragraph is ordinary again.")
mark(90, "The next paragraph is ordinary again.", 2)
Y[0] += 10; ln(72, "Notes", "bd", 12)
ln(72, "1. First endnote.")
ln(72, "2. Second endnote that runs over two lines of the note area")
ln(84, "and finishes here.")
Y[0] += 10; ln(72, "References", "bd", 12)
ln(72, "Acton, Thomas. 1974. Gypsy Politics and Social Change. London: Routledge &")
ln(90, "Kegan Paul.")
ln(72, "Fraser, Angus. 1992. The Gypsies. Oxford: Blackwell.")
ep = os.path.join(d, "end.pdf"); doc.save(ep)
r = subprocess.run([sys.executable, EX, ep, "-o", os.path.join(d, "e.md")], capture_output=True, text=True)
e = open(os.path.join(d, "e.md"), encoding="utf-8").read()
t("endnotes section parsed as notes, heading dropped", "[^2]: Second endnote that runs over two lines of the note area and finishes here." in e and "Notes" not in e and "EXTRACT OK" in r.stdout, e + r.stdout)
t("same-size indented quotation -> one > paragraph", "> An indented quotation set at the body size runs across two lines between wider margins and ends." in e, e)
t("first-line-indented paragraphs not mistaken for quotations", "\nA first paragraph opens" in "\n" + e and "\nThe next paragraph is ordinary again.[^2]" in "\n" + e, e)
bibt = open(os.path.join(d, "e_bib.txt"), encoding="utf-8").read().splitlines() if os.path.exists(os.path.join(d, "e_bib.txt")) else []
t("reference list with hanging indent -> one entry per line in _bib.txt", bibt == ["Acton, Thomas. 1974. Gypsy Politics and Social Change. London: Routledge & Kegan Paul.", "Fraser, Angus. 1992. The Gypsies. Oxford: Blackwell."], bibt)

# ---- hand-set page 3: raised note numbers; a continuation line that starts with "2 marca"
doc = pymupdf.open(); pg = doc.new_page()
pg.insert_font(fontname="rg", fontbuffer=FONTS["rg"])
Y = [90]
ln(72, "Body text with the first marker and more words")
mark(72, "Body text with the first marker and more words", 1)
ln(72, "and a second marker at the end of this line")
mark(72, "and a second marker at the end of this line", 2)
for _ in range(8):
    ln(72, "Further body text fills the page so that the body size dominates the page.")
for y, x, txt, fs in [(700, 72, "1", 6), (704, 77, "First note that runs on and on until the end of the line, then", 9),
                      (716, 72, "2 marca 1937 r. it continues on this line.", 9), (730, 72, "2", 6), (734, 77, "Second note.", 9)]:
    pg.insert_text((x, y), txt, fontname="rg", fontsize=fs)
cp = os.path.join(d, "cont.pdf"); doc.save(cp)
r = subprocess.run([sys.executable, EX, cp, "-o", os.path.join(d, "c.md")], capture_output=True, text=True)
cm_ = open(os.path.join(d, "c.md"), encoding="utf-8").read()
t("continuation line starting with a number is not taken for the next note (raised numbers)", "EXTRACT OK" in r.stdout and "then 2 marca 1937 r. it continues on this line." in cm_ and "[^2]: Second note." in cm_, r.stdout + cm_)

n, ok = len(res), sum(res)
print(f"PDF ALL PASS {n}/{n}" if ok == n else f"PDF FAILED {n - ok}/{n}")
