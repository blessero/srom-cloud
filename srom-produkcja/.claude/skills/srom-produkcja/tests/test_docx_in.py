"""G13: DOCX import — footnotes kept and moved under their paragraph, italics and small caps kept,
tracked changes/comments counted, hand-typed superscript numbers, manual headings and numbered
lists flagged; the chain import -> normalize -> build stops on the forbidden 'Tamże'."""
import os, re, sys, json, subprocess, tempfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
S = os.path.join(ROOT, "scripts")
d = tempfile.mkdtemp()
res = []
def t(name, ok, detail=""):
    res.append(ok); print(("PASS " if ok else "FAIL ") + name + ("" if ok else "\n   " + str(detail)[:1500]))
docx = os.path.join(d, "in.docx")
subprocess.run(["pandoc", os.path.join(ROOT, "tests", "fixtures", "docx_in_src.md"), "-o", docx], check=True)
r = subprocess.run([sys.executable, os.path.join(S, "docx_in.py"), docx, "-o", os.path.join(d, "a.md")], capture_output=True, text=True)
md = open(os.path.join(d, "a.md"), encoding="utf-8").read()
rep = open(os.path.join(d, "a_import.md"), encoding="utf-8").read()
blocks = [b for b in re.split(r"\n\s*\n", md) if b.strip()]
t("footnote definitions directly under their paragraph (incl. 2nd paragraph of note 2)",
  blocks[1].startswith("Pierwszy akapit") and blocks[2].startswith("[^1]:") and blocks[3].startswith("[^2]:") and blocks[4].startswith("    Drugi akapit"), blocks[:6])
t("italics and small caps preserved", "*kursywą*" in md and "[Ficowski]{.smallcaps}" in md and "*Cyganie na polskich drogach*" in md, md)
t("tracked change accepted and counted, comment dropped and counted", "wstawką" in md and "1 insertions" in rep and "comments dropped: 1" in rep, rep)
t("hand-typed superscript digit flagged as possible fake footnote", "FAKE-NOTE" in rep, rep)
t("bold-only paragraph flagged as manual heading", "BOLDPARA" in rep, rep)
t("numbered list kept as a list (kanon allows numbered lists)", "1.  pierwszy punkt" in md or "1. pierwszy punkt" in md, md)
t("link reduced to text, URL listed", "oraz link." in md and "https://example.org" in rep, md)
n_md = os.path.join(d, "n.md")
subprocess.run([sys.executable, os.path.join(S, "normalize.py"), os.path.join(d, "a.md"), "-o", n_md, "--log", os.path.join(d, "n.log")], capture_output=True)
b = subprocess.run([sys.executable, os.path.join(S, "build.py"), n_md, "--out", os.path.join(d, "b")], capture_output=True, text=True)
brep = open(os.path.join(d, "b", "n_report.md"), encoding="utf-8").read()
t("import -> normalize -> build: unkeyed 'Tamże' blocks the build (kanon §7.3)", b.returncode == 1 and "kanon linter:" in brep and "ABBR" in brep, brep[:1500])
# ---- Zotero citation fields are harvested; the author's reference list is split off to _bib.txt
import docx as _docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
D = _docx.Document()
D.add_heading("1. Wstęp", level=1)
p = D.add_paragraph("Tekst przed cytowaniem ")
cit = {"citationID": "abc", "properties": {"formattedCitation": "(Ficowski 1985, 15)", "plainCitation": "(Ficowski 1985, 15)", "noteIndex": 0},
       "citationItems": [{"id": 123, "uris": ["http://zotero.org/users/1/items/ABCD1234"], "locator": "15", "label": "page",
                          "itemData": {"id": 123, "type": "book", "title": "Cyganie na polskich drogach", "publisher": "Wydawnictwo Literackie",
                                       "publisher-place": "Kraków", "author": [{"family": "Ficowski", "given": "Jerzy"}], "issued": {"date-parts": [["1985"]]},
                                       "ISBN": "83-08-01234-5"}}],
       "schema": "https://github.com/citation-style-language/schema/raw/master/csl-citation.json"}
def fld(par, instr, shown):
    for kind, txt in (("begin", None), ("instr", instr), ("separate", None), ("text", shown), ("end", None)):
        r = par.add_run(txt if kind == "text" else None)
        if kind == "instr":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
        elif kind != "text":
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r._r.append(e)
fld(p, " ADDIN ZOTERO_ITEM CSL_CITATION " + json.dumps(cit) + " ", "(Ficowski 1985, 15)")
p.add_run(" i dalej.")
D.add_heading("Bibliografia", level=1)
D.add_paragraph("Ficowski J., Cyganie na polskich drogach, Kraków 1985.")
D.add_paragraph("Mróz L., Dzieje Cyganów-Romów w Rzeczypospolitej, Warszawa 2001.")
zd = os.path.join(d, "z.docx"); D.save(zd)
r = subprocess.run([sys.executable, os.path.join(S, "docx_in.py"), zd, "-o", os.path.join(d, "z.md")], capture_output=True, text=True)
zrep = open(os.path.join(d, "z_import.md"), encoding="utf-8").read()
cited = open(os.path.join(d, "z_cited.md"), encoding="utf-8").read() if os.path.exists(os.path.join(d, "z_cited.md")) else ""
zrefs = json.load(open(os.path.join(d, "z_refs.json"), encoding="utf-8")) if os.path.exists(os.path.join(d, "z_refs.json")) else []
t("Zotero field harvested: citation keyed surnameYEAR with Polish locator", "[@ficowski1985, s. 15]" in cited and [x["id"] for x in zrefs] == ["ficowski1985"], zrep + cited)
t("author's reference list removed from the text and written to _bib.txt", "Bibliografia" not in cited and open(os.path.join(d, "z_bib.txt"), encoding="utf-8").read().count("\n") == 2, cited)
b = subprocess.run([sys.executable, os.path.join(S, "build.py"), os.path.join(d, "z_cited.md"), "--refs", os.path.join(d, "z_refs.json"), "--out", os.path.join(d, "zb")], capture_output=True, text=True)
zbt = open(os.path.join(d, "zb", "z_cited.txt"), encoding="utf-8").read() if os.path.exists(os.path.join(d, "zb", "z_cited.txt")) else ""
t("harvested article builds; the author-date citation becomes a SROM footnote", b.returncode == 0 and "cytowaniem[1] i dalej" in zbt and "J. Ficowski, Cyganie na polskich drogach, Wydawnictwo Literackie, Kraków 1985, s. 15." in zbt, b.stdout + zbt)

D2 = _docx.Document(); D2.add_paragraph("(1)\tMe\tdikhav"); D2.add_paragraph("Zwykły akapit."); td = os.path.join(d, "tabs.docx"); D2.save(td)
r = subprocess.run([sys.executable, os.path.join(S, "docx_in.py"), td, "-o", os.path.join(d, "tabs.md")], capture_output=True, text=True)
trep = open(os.path.join(d, "tabs_import.md"), encoding="utf-8").read()
t("tab-aligned paragraph (interlinear example typed in Word) flagged", "TABS" in trep and "(1)" in trep and trep.count("TABS") == 1 and "IMPORT CHECK" in r.stdout, trep)

# ---- E9 --typed-notes: a file laid out like its PDF — superscript digits in the body, each page's notes typed as
# numbered paragraphs after that page's text, paragraphs split at page breaks (Dom file, vol. 18)
from docx.enum.style import WD_STYLE_TYPE
D = _docx.Document()
D.styles.add_style("Note", WD_STYLE_TYPE.PARAGRAPH)
def para(style, *parts):
    p = D.add_paragraph(style=style)
    for x in parts:
        r = p.add_run(x.lstrip("^/"))
        r.font.superscript = x.startswith("^")
        r.italic = x.startswith("/")
BT, NT = "Body Text", "Note"
para(BT, "Pierwszy akapit z przypisem", "^1", " i drugim", "^2", " oraz trzecim", "^3", " ciągnie się na następną")
para(NT, "^1", " Pierwszy przypis, ", "/Tytuł", ".")
para(NT, "2 Drugi przypis. ", "^3", " Trzeci przypis wpisany w drugi,")
para(NT, "ciąg dalszy trzeciego.")
para(BT, "900430271992")
para(BT, "stronę i tu się kończy.")
para(BT, "Nowy akapit", "^4", " bez kropki")
para(BT, "5. Dalszy tekst po zgubionym odsyłaczu", "^6", ".")
para(NT, "^4", " Czwarty.")
para(NT, "^5", " Piąty.")
para(NT, "^6", " Szósty.")
para(BT, "Akapit z odsyłaczem", "^7", " urwany na")
para(NT, "^7", " Siódmy.")
para(BT, "zdaniu, dalej", "^10", ".")
para(NT, "^10", " Dziesiąty.")
para(BT, "Odsyłacz bez przypisu", "^11", ".")
tn = os.path.join(d, "typed.docx"); D.save(tn)
r = subprocess.run([sys.executable, os.path.join(S, "docx_in.py"), tn, "-o", os.path.join(d, "tn.md"), "--typed-notes"], capture_output=True, text=True)
md = open(os.path.join(d, "tn.md"), encoding="utf-8").read()
trep = open(os.path.join(d, "tn_import.md"), encoding="utf-8").read()
defs = dict(re.findall(r"^\[\^(\d+)\]: (.*)$", md, re.M))
t("E9: typed notes become real notes, labelled with the source numbers (8 = 1–7, 10)",
  sorted(map(int, defs)) == [1, 2, 3, 4, 5, 6, 7, 10] and "8 markers paired" in trep, (r.stdout, r.stderr, md, trep))
t("E9: note text kept with italics; the number is not part of it", defs.get("1") == "Pierwszy przypis, *Tytuł*.", defs)
t("E9: a note typed inside the previous one is split off (Dom: note 2)",
  defs.get("2") == "Drugi przypis." and defs.get("3", "").startswith("Trzeci przypis"), defs)
t("E9: a note's continuation line (note style) is joined to it",
  defs.get("3") == "Trzeci przypis wpisany w drugi, ciąg dalszy trzeciego.", defs)
t("E9: a body paragraph split by the page (and its notes) is joined; page-ID string dropped",
  "na następną stronę i tu się kończy." in md and "900430271992" not in md and "page-ID string dropped" in trep, md)
t("E9: marker typed at a paragraph start („5. Dalszy…”) repaired, listed for verification",
  "bez kropki[^5]. Dalszy tekst" in md and "REPAIRED: marker 5" in trep, (md, trep))
t("E9: never joined across a missing page — note and marker numbers skip (7 → 10): listed, not guessed",
  "urwany na\n" in md and "\nzdaniu, dalej[^10]" in md and "LOST TEXT?" in trep and "GAP: notes 8–9" in trep, (md, trep))
t("E9: a marker without a note stays text and is listed; the import says IMPORT CHECK",
  "[^11]" not in md and "MARKER WITHOUT NOTE: [11]" in trep and "IMPORT CHECK" in r.stdout, (md, trep, r.stdout))
c = subprocess.run([sys.executable, os.path.join(S, "check.py"), os.path.join(d, "tn.md")], capture_output=True, text=True)
t("E9: the result passes the integrity check", c.returncode == 0, c.stdout)
r0 = subprocess.run([sys.executable, os.path.join(S, "docx_in.py"), tn, "-o", os.path.join(d, "tn0.md")], capture_output=True, text=True)
t("E9: without --typed-notes nothing is converted (0 notes, typed digits flagged)",
  "[^" not in open(os.path.join(d, "tn0.md"), encoding="utf-8").read() and "FAKE-NOTE" in open(os.path.join(d, "tn0_import.md"), encoding="utf-8").read())
rx = subprocess.run([sys.executable, os.path.join(S, "docx_in.py"), docx, "-o", os.path.join(d, "mixed.md"), "--typed-notes"], capture_output=True, text=True)
t("E9: a file that already has real Word notes is refused (mixed notes)", rx.returncode != 0 and "already has real Word footnotes" in rx.stderr, rx.stderr)

n, ok = len(res), sum(res)
print(f"DOCXIN ALL PASS {n}/{n}" if ok == n else f"DOCXIN FAILED {n - ok}/{n}")
