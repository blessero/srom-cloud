#!/usr/bin/env python3
"""
build.py — SROM-MD (+ CSL-JSON refs) -> InDesign-ready DOCX, with verification.

    python3 build.py article.md --refs refs.json --out build/ [--config config/styles.json] [--draft]

Outputs in --out:
    <stem>.docx                 import into InDesign with the "SROM – pandoc" Word import preset
    <stem>_report.md            what was built, warnings, flags, verification results
    <stem>_postimport.jsx       run right after placing the DOCX: styles/overrides/footnote count
    <stem>_ibidem.jsx           run after final layout: Ibidem on the same column as the preceding note (§7.3)
    <stem>_gwiazdki.jsx         (only with non-author notes) after layout: asterisks per page for the * series (§7.1)
    <stem>.txt                  plain text (for linting / proofreading diff)

Exit code 1 = verification failed (do not typeset). Placeholders [BRAK …] fail the build unless --draft.
"""
import argparse, copy, json, os, re, shutil, subprocess, sys, tempfile, zipfile
from collections import Counter, OrderedDict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSL = os.path.join(ROOT, "csl", "srom.csl")
LUA = os.path.join(ROOT, "lua", "srom_post.lua")
JSX_TPL = os.path.join(ROOT, "indesign", "srom_ibidem_check.jsx.tpl")
POSTIMPORT_TPL = os.path.join(ROOT, "indesign", "srom_postimport.jsx.tpl")
ASTERISK_TPL = os.path.join(ROOT, "indesign", "srom_gwiazdki.jsx.tpl")
FROM = ("markdown-smart-superscript-subscript-strikeout-raw_html-raw_tex-tex_math_dollars"
        "-implicit_figures-fancy_lists-example_lists-task_lists-auto_identifiers")
SECTIONS = OrderedDict([("I", "Wykaz skrótów"), ("II", "Źródła archiwalne"), ("III", "Źródła terenowe"),
                        ("IV", "Źródła drukowane i prawne"), ("V", "Źródła internetowe"), ("VI", "Literatura przedmiotu")])
ROMAN = list(SECTIONS.keys())
BIB_NUMBERS = [False]   # set from config "bibliography.section_numbers"
ITALIC_STYLE = ["Kursywa"]  # set from config in main()
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


PANDOC_VERSION = subprocess.run(["pandoc", "--version"], capture_output=True, text=True).stdout.split("\n")[0].replace("pandoc ", "")


def run(cmd, **kw):
    r = subprocess.run(cmd, capture_output=True, text=True, **kw)
    return r.returncode, r.stdout, r.stderr


def pandoc_json(md_path, extra=()):
    code, out, err = run(["pandoc", md_path, "-f", FROM, "-t", "json", *extra])
    if code:
        sys.exit(f"pandoc failed:\n{err}")
    return json.loads(out), err


def walk(x, fn):
    if isinstance(x, list):
        for i in x:
            walk(i, fn)
    elif isinstance(x, dict):
        fn(x)
        if "c" in x:
            walk(x["c"], fn)


def ser(x):
    """AST -> readable text with *italics* (for reports and the ibidem map)."""
    if isinstance(x, list):
        return "".join(ser(i) for i in x)
    t, c = x["t"], x.get("c")
    if t == "Str": return c
    if t in ("Space", "SoftBreak"): return " "
    if t == "Emph": return "*" + ser(c) + "*"
    if t == "Quoted": return ("„" + ser(c[1]) + "”") if c[0]["t"] == "DoubleQuote" else ("»" + ser(c[1]) + "«")
    if t == "Span":
        kv = dict(c[0][2])
        return ("*" + ser(c[1]) + "*") if kv.get("custom-style") == ITALIC_STYLE[0] else ser(c[1])
    if t == "Link": return ser(c[1])
    if t == "SmallCaps": return ser(c)
    if t == "Cite": return ser(c[1])
    if t == "Note": return ""
    if t in ("Para", "Plain"): return ser(c)
    if t in ("Strong", "Underline", "Superscript", "Subscript", "Strikeout"): return ser(c)
    return ""


def pre_notes(doc):
    """note sequence BEFORE citeproc, in footnote-number order: every Note, and every Cite in the main
    text (citeproc turns it into a note). Each entry: literal text outside citations."""
    out = []

    def rec(x, in_note):
        if isinstance(x, list):
            for i in x:
                rec(i, in_note)
        elif isinstance(x, dict):
            t = x.get("t")
            if t == "Note" and not in_note:
                out.append(literal_text(x["c"]))
                return
            if t == "Cite" and not in_note:
                out.append("")
                return
            if "c" in x:
                rec(x["c"], in_note)
    rec(doc["blocks"], False)
    return out


def literal_text(x):
    """text of a note outside its citations"""
    if isinstance(x, list):
        return "".join(literal_text(i) for i in x)
    if not isinstance(x, dict):
        return ""
    t, c = x.get("t"), x.get("c")
    if t == "Cite":
        return " "
    if t == "Str":
        return c
    if t in ("Space", "SoftBreak"):
        return " "
    return literal_text(c) if isinstance(c, list) else ""


SEE_WORDS = re.compile(r"(?i)\b(zob\.|por\.|np\.|też|także|szerzej|więcej|cyt\.\s+za|za:)")


def extra_source(literal):
    """does the literal part of a note look like another source (archival unit, press issue, act)?"""
    rest = SEE_WORDS.sub(" ", literal)
    return bool(re.search(r"\d", rest))


def note_dicts(doc):
    out = []
    walk(doc["blocks"], lambda x: out.append(x) if x.get("t") == "Note" else None)
    return out


def notes_text(doc):
    out = []
    def fn(x):
        if x.get("t") == "Note":
            out.append(" ".join(ser(b) for b in x["c"]))
    walk(doc["blocks"], fn)
    return out


# ------------------------------------------------------------------ reference.docx from config
def make_reference_docx(cfg, path):
    import docx
    from docx.enum.style import WD_STYLE_TYPE
    tmp = path + ".base.docx"
    with open(tmp, "wb") as f:
        f.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True).stdout)
    d = docx.Document(tmp)
    names = {s.name for s in d.styles}
    for key, name in cfg["paragraph"].items():
        if key == "footnote" or name in names:
            continue
        st = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = d.styles["Normal"]
        if key in ("h1", "h2", "bib_title", "bib_section"):
            st.font.bold = True
        if key == "quote":
            st.paragraph_format.left_indent = docx.shared.Cm(1)
    for key, name in cfg["character"].items():
        if name in names:
            continue
        st = d.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
        if key == "italic":
            st.font.italic = True
        if key == "smallcaps":
            st.font.small_caps = True
        if key == "asterisk_ref":
            st.font.superscript = True
    d.save(path)
    os.remove(tmp)


# ------------------------------------------------------------------ bibliography assembly
BIB_BLOCK = re.compile(r"^::: *\{#bibliografia\} *\n(.*?)^:::[ \t]*$", re.S | re.M)


def _norm(s):
    import unicodedata
    s = s.replace("ł", "l").replace("Ł", "L")
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn").lower().strip(" .")


def parse_literal_bib(block, report=None):
    """Literal sections inside the ::: {#bibliografia} block -> {section id: [entries]}.
    A section heading is its name ("## Źródła archiwalne"); an old-style numeral ("## II. …") is accepted."""
    by_name = {_norm(v): k for k, v in SECTIONS.items()}
    secs, cur = OrderedDict(), None
    for ln in block.split("\n"):
        m = re.match(r"^##\s+(?:(?:I|II|III|IV|V|VI)\.\s+)?(.+?)\s*$", ln)
        if m:
            cur = by_name.get(_norm(m.group(1)))
            if cur is None:
                if report is not None:
                    report["errors"].append(f"unknown bibliography section '{m.group(1)}' — use one of: {', '.join(SECTIONS.values())}")
                cur = "VI"
            secs.setdefault(cur, [])
        elif ln.startswith("# "):
            continue
        elif ln.strip() and cur:
            secs[cur].append(ln.strip())
        elif ln.strip():
            secs.setdefault("VI", []).append(ln.strip())
    return secs


def csl_entries(keys, refs_path, workdir):
    if not keys:
        return []
    mini = os.path.join(workdir, "bib_mini.md")
    with open(mini, "w", encoding="utf-8") as f:
        f.write("---\nlang: pl-PL\nnocite: '" + ", ".join("@" + k for k in keys) + "'\n---\n\n")
    code, out, err = run(["pandoc", mini, "-f", FROM, "--citeproc", "--csl", CSL, "--bibliography", refs_path,
                          "-t", "markdown-smart-citations", "--wrap=none"])
    if code:
        sys.exit(err)
    entries, cur = [], None
    for ln in out.split("\n"):
        if ln.startswith("::: {#ref-"):
            cur = []
        elif ln.startswith(":::"):
            if cur is not None:
                entries.append(" ".join(cur))
            cur = None
        elif cur is not None and ln.strip():
            cur.append(ln.strip())
    return entries


def section_of(ref):
    s = ref.get("srom-section")
    if s:
        return s
    return "V" if ref.get("type") in ("webpage", "post", "post-weblog") and not ref.get("DOI") else "VI"


def unsmallcap_literals(entries, keys, refs):
    """institutional authors (CSL 'literal' names) are not surnames: no small caps (kanon §9.3)"""
    lits = {p["literal"] for k in keys for role in ("author", "editor") for p in refs[k].get(role, []) if "literal" in p}
    out = []
    for e in entries:
        for lit in lits:
            e = e.replace("[" + lit + "]{.smallcaps}", lit)
        out.append(e)
    return out


def build_bibliography(md_text, refs, cited, refs_path, workdir, report):
    m = BIB_BLOCK.search(md_text)
    literal = parse_literal_bib(m.group(1), report) if m else OrderedDict()
    gen = OrderedDict()
    for k in cited:
        gen.setdefault(section_of(refs[k]), []).append(k)
    present = [r for r in ROMAN if literal.get(r) or gen.get(r)]
    lines = ["::: {.bibliografia}", "# Bibliografia", ""]
    for i, r in enumerate(present):
        if literal.get(r) and gen.get(r):
            report["errors"].append(f"Bibliography section {r} has both literal entries and CSL entries — keep one source per section")
        if len(present) > 1:
            # kanon: only the sections an article uses are printed, without numerals (config can restore them)
            label = f"{ROMAN[i]}. {SECTIONS[r]}" if BIB_NUMBERS[0] else SECTIONS[r]
            lines += [f"## {label}", ""]
        ents = literal.get(r) or unsmallcap_literals(csl_entries(gen.get(r, []), refs_path, workdir), gen.get(r, []), refs)
        for e in ents:
            lines += [e, ""]
    lines.append(":::")
    block = "\n".join(lines)
    report["bib_sections"] = {r: len(literal.get(r) or gen.get(r) or []) for r in present}
    return (md_text[:m.start()] + block + md_text[m.end():]) if m else (md_text.rstrip() + "\n\n" + block + "\n")


# ------------------------------------------------------------------ DOCX post-processing + verification
def postprocess_docx(path, cfg, report):
    from lxml import etree
    tmpdir = tempfile.mkdtemp()
    with zipfile.ZipFile(path) as z:
        z.extractall(tmpdir)
        names = z.namelist()

    def load(n):
        return etree.parse(os.path.join(tmpdir, n))

    def save(t, n):
        t.write(os.path.join(tmpdir, n), xml_declaration=True, encoding="UTF-8", standalone=True)

    # styles: rename footnote style, drop language defaults, map ids -> names
    st = load("word/styles.xml")
    id2name = {}
    for s in st.getroot().iter(W + "style"):
        sid = s.get(W + "styleId")
        nm = s.find(W + "name")
        if sid == "FootnoteText" and nm is not None:
            nm.set(W + "val", cfg["paragraph"]["footnote"])
        if nm is not None:
            id2name[sid] = nm.get(W + "val")
    for lang in list(st.getroot().iter(W + "lang")):
        lang.getparent().remove(lang)
    save(st, "word/styles.xml")

    fixed_refs = 0
    for part in ("word/document.xml", "word/footnotes.xml"):
        if part not in names:
            continue
        t = load(part)
        root = t.getroot()
        # footnote reference runs: drop Word's FootnoteReference style (InDesign's Footnote Options format the number)
        for tag in ("footnoteReference", "footnoteRef"):
            for ref in root.iter(W + tag):
                run_el = ref.getparent()
                rpr = run_el.find(W + "rPr")
                if rpr is not None:
                    run_el.remove(rpr)
                    fixed_refs += 1
                if tag == "footnoteRef":
                    nxt = run_el.getnext()
                    tt = nxt.find(W + "t") if nxt is not None else None
                    if tt is not None and tt.text == " " and len(nxt) == 1:
                        nxt.getparent().remove(nxt)
                    elif tt is not None and tt.text and tt.text.startswith(" "):
                        tt.text = tt.text.lstrip(" ")
        save(t, part)
    report["footnote_ref_styles_removed"] = fixed_refs
    # literal tabs (interlinear examples) must be <w:tab/> elements, not characters inside <w:t>
    for part in ("word/document.xml", "word/footnotes.xml"):
        if part not in names:
            continue
        t = load(part)
        changed = False
        for wt_el in list(t.getroot().iter(W + "t")):
            if wt_el.text and "\t" in wt_el.text:
                run_el = wt_el.getparent()
                idx = list(run_el).index(wt_el)
                pieces = wt_el.text.split("\t")
                run_el.remove(wt_el)
                new_els = []
                for j, piece in enumerate(pieces):
                    if j:
                        new_els.append(etree.Element(W + "tab"))
                    if piece:
                        e = etree.Element(W + "t")
                        e.text = piece
                        e.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                        new_els.append(e)
                for off, e in enumerate(new_els):
                    run_el.insert(idx + off, e)
                changed = True
        if changed:
            save(t, part)

    out = path + ".tmp"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.write(os.path.join(tmpdir, n), n)
    shutil.move(out, path)
    shutil.rmtree(tmpdir)
    return id2name


def verify_docx(path, cfg, id2name, expected_notes, report, expected_ast=(0, 0)):
    """expected_ast: (non-author notes, title notes) — the asterisk series (kanon § 7.1)"""
    from lxml import etree
    z = zipfile.ZipFile(path)
    allowed_p = set(cfg["paragraph"].values())
    allowed_c = set(cfg["character"].values())
    doc = etree.fromstring(z.read("word/document.xml"))
    fn = etree.fromstring(z.read("word/footnotes.xml")) if "word/footnotes.xml" in z.namelist() else None
    res = []

    def styles_used(root, kind):
        c = Counter()
        for el in root.iter(W + kind):
            c[id2name.get(el.get(W + "val"), el.get(W + "val"))] += 1
        return c

    pstyles = styles_used(doc, "pStyle")
    unstyled = sum(1 for p in doc.iter(W + "p") if p.find(W + "pPr/" + W + "pStyle") is None and "".join(p.itertext()).strip())
    bad_p = {k: v for k, v in pstyles.items() if k not in allowed_p}
    res.append(("paragraph styles ⊆ config", not bad_p and unstyled == 0,
                f"used {dict(pstyles)}; foreign {bad_p}; unstyled {unstyled}"))
    cstyles = styles_used(doc, "rStyle") + (styles_used(fn, "rStyle") if fn is not None else Counter())
    bad_c = {k: v for k, v in cstyles.items() if k not in allowed_c}
    res.append(("character styles ⊆ config", not bad_c, f"used {dict(cstyles)}; foreign {bad_c}"))
    direct = 0
    for root in (doc, fn):
        if root is None:
            continue
        for rpr in root.iter(W + "rPr"):
            for tag in ("i", "b", "smallCaps", "caps", "u", "iCs", "bCs"):
                if rpr.find(W + tag) is not None:
                    direct += 1
    res.append(("no direct italic/bold/caps runs", direct == 0, f"{direct} direct-formatted runs"))
    notes = [f for f in fn.iter(W + "footnote") if f.get(W + "type") is None] if fn is not None else []
    res.append(("footnote count", len(notes) == expected_notes, f"docx {len(notes)} / source {expected_notes}"))
    if fn is not None:
        fpstyles = styles_used(fn, "pStyle")
        ok = set(fpstyles) <= {cfg["paragraph"]["footnote"]}
        res.append(("footnote paragraphs use footnote style only", ok, str(dict(fpstyles))))
        lead = 0
        for f in notes:
            p = f.find(W + "p")
            txt = "".join(t.text or "" for t in p.iter(W + "t")) if p is not None else ""
            if txt.startswith(" "):
                lead += 1
        res.append(("no leading space in notes", lead == 0, f"{lead} notes start with a space"))
        empty_first = sum(1 for f in notes if not "".join(t.text or "" for t in f.find(W + "p").iter(W + "t")).strip())
        res.append(("footnote number not in an empty paragraph", empty_first == 0, f"{empty_first}"))
        sys.path.insert(0, os.path.join(ROOT, "scripts"))
        import check as integrity
        na_in = [i for i, f in enumerate(notes, 1)
                 if integrity.na_kind("".join(t.text or "" for t in f.iter(W + "t")))]
        res.append(("no translator/editorial note among the numbered footnotes (§7.1)", not na_in, f"footnotes {na_in}"))
    n_na, n_title = expected_ast
    ast_ref, ast_note = cfg["character"]["asterisk_ref"], cfg["paragraph"]["asterisk_note"]
    marks = sum(1 for r in doc.iter(W + "rStyle") if id2name.get(r.get(W + "val"), r.get(W + "val")) == ast_ref)
    heads = 0
    for p in doc.iter(W + "p"):
        ps = p.find(W + "pPr/" + W + "pStyle")
        if ps is not None and id2name.get(ps.get(W + "val"), ps.get(W + "val")) == ast_note \
                and "".join(t.text or "" for t in p.iter(W + "t")).startswith("* "):
            heads += 1
    res.append(("asterisk markers in the text = translator/editorial notes", marks == n_na, f"markers {marks} / notes {n_na}"))
    res.append(("asterisk notes at the end = notes + title note", heads == n_na + n_title,
                f"'{ast_note}' notes {heads} / expected {n_na} + {n_title}"))
    hl = len(list(doc.iter(W + "hyperlink"))) + (len(list(fn.iter(W + "hyperlink"))) if fn is not None else 0)
    res.append(("no hyperlinks", hl == 0, f"{hl}"))
    def wt(root):
        return "\n".join("".join(t.text or "" for t in p.iter(W + "t")) for p in root.iter(W + "p")) if root is not None else ""
    alltext = wt(doc) + "\n" + wt(fn)
    for label, rx in [("no straight double quotes", r'"'), ("no em dash", "—"), ("no double spaces", "  "),
                      ("no ASCII ellipsis", r"\.\.\."), ("no English quotes “", "“")]:
        n = len(re.findall(rx, alltext))
        res.append((label, n == 0, f"{n}"))
    report["verify"] = res
    return all(ok for _, ok, _ in res)


# ------------------------------------------------------------------ ibidem map (§7.3: Ibidem only on the same column)
def noibid_csl(workdir):
    s = open(CSL, encoding="utf-8").read()
    s = s.replace('<if position="ibid-with-locator">', '<if position="ibid-with-locator" variable="srom-never" match="all">')
    s = s.replace('<else-if position="ibid">', '<else-if position="ibid" variable="srom-never" match="all">')
    p = os.path.join(workdir, "noibid.csl")
    open(p, "w", encoding="utf-8").write(s)
    return p


def placeholders_by_key(refs_path, keys, workdir):
    """render each cited work once (bibliography form) and collect its [BRAK …] placeholders"""
    if not keys:
        return {}
    mini = os.path.join(workdir, "brak.md")
    open(mini, "w", encoding="utf-8").write("---\nnocite: '" + ", ".join("@" + k for k in keys) + "'\n---\n\nX\n")
    doc, _ = pandoc_json(mini, ["--citeproc", "--csl", CSL, "--bibliography", refs_path, "-M", "lang=pl-PL"])
    out = {}
    def fn(x):
        if x.get("t") == "Div" and x["c"][0][0].startswith("ref-"):
            txt = ser(x["c"][1])
            found = re.findall(r"\[BRAK [A-ZĄĆĘŁŃÓŚŹŻ ]+\]", txt)
            if found:
                out[x["c"][0][0][4:]] = sorted(set(found))
    walk(doc["blocks"], fn)
    return out


def write_queries(outdir, stem, nopage, notes, refs, cited, cite_count, refs_path, workdir, report, extra=(), label_map=None):
    """Author/editor query sheet: everything a human must answer or confirm, one row each.
    <stem>_pytania.md (readable) and <stem>_pytania.csv (Excel, ';', UTF-8 BOM)."""
    rows = []   # (adresat, rodzaj, przypis, dzieło, szczegół)
    def cut(t, n=110):
        return t if len(t) <= n else t[:n] + "…"
    for no, kind, keys, ctx in nopage:
        note = notes[no - 1] if isinstance(no, int) and 0 < no <= len(notes) else ""
        if kind == "quote":
            rows.append(("autor", "cytat bez numeru strony — prosimy o stronę", str(no), keys, f"…{ctx[-80:]} | przypis: {cut(note)}"))
        else:
            rows.append(("redakcja", "odwołanie do całości dzieła (bez strony) — potwierdzić", str(no), keys, f"…{ctx[-80:]} | przypis: {cut(note)}"))
    for key, found in sorted(placeholders_by_key(refs_path, cited, workdir).items()):
        rows.append(("autor", "brak danych bibliograficznych: " + ", ".join(found), "", key, cut(refs[key].get("title", ""))))
    for k in cited:
        r = refs[k]
        yr = (r.get("issued", {}).get("date-parts") or [[None]])[0][0]
        if r.get("type") == "book" and not r.get("ISBN") and isinstance(yr, int) and yr >= 1970:
            rows.append(("autor", "monografia po 1970 r. bez ISBN (§9.1) — podać lub potwierdzić brak", "", k, cut(r.get("title", ""))))
        if cite_count[k] >= 2 and not r.get("title-short") and len((r.get("title") or "").split()) > 5:
            rows.append(("redakcja", "długi tytuł bez formy skróconej (title-short) — ustalić (§7.3)", "", k, cut(r.get("title", ""))))
    import csv
    for path in extra:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rd = csv.reader(fh, delimiter=";")
            head = next(rd, None)
            if head and [h.strip() for h in head[:5]] != ["adresat", "rodzaj", "przypis", "dzieło", "szczegóły"]:
                report["errors"].append(f"{os.path.basename(path)}: query sheet header must be adresat;rodzaj;przypis;dzieło;szczegóły")
                continue
            for row in rd:
                if any(c.strip() for c in row):
                    row = (row + [""] * 5)[:5]
                    lab = row[2].strip().strip("[]^")
                    if label_map and lab in label_map:
                        row[2] = str(label_map[lab])     # source note label -> printed number (E6)
                    rows.append(tuple(row))
    md = [f"# Pytania i potwierdzenia — {stem}", "",
          "| adresat | rodzaj | przypis | dzieło | szczegóły |", "|---|---|---|---|---|"]
    md += [f"| {a} | {b} | {c} | {d} | {e.replace('|', '/')} |" for a, b, c, d, e in rows] or ["| — | brak pytań | | | |"]
    open(os.path.join(outdir, stem + "_pytania.md"), "w", encoding="utf-8").write("\n".join(md) + "\n")
    with open(os.path.join(outdir, stem + "_pytania.csv"), "w", encoding="utf-8-sig", newline="") as fh:
        wr = csv.writer(fh, delimiter=";")
        wr.writerow(["adresat", "rodzaj", "przypis", "dzieło", "szczegóły"])
        wr.writerows(rows)
    q = Counter(r[1].split(" — ")[0].split(":")[0] for r in rows)
    report["queries"] = rows
    if rows:
        report["warnings"].append(f"query sheet {stem}_pytania.md: " + ", ".join(f"{v}× {k}" for k, v in q.items()))


def write_jsx(rows, total, cfg, outdir, stem, ast_texts=(), title=False):
    """Per-article InDesign scripts (report-only): Ibidem column check + post-import check; with non-author
    notes also the asterisk-series list (ast_texts: the notes in order, each "* …", title note first if title)."""
    def js(x):
        return json.dumps(x, ensure_ascii=True)
    body = ",\n".join(f"  {i}: {{ibid: {js(x)}, full: {js(y)}}}" for i, x, y in rows)
    tpl = open(JSX_TPL, encoding="utf-8").read()
    open(os.path.join(outdir, stem + "_ibidem.jsx"), "w", encoding="utf-8").write(
        tpl.replace("/*MAP*/", body).replace("/*TOTAL*/", str(total)))
    extra = cfg.get("template_extra", {})
    para = [v for k, v in cfg["paragraph"].items()] + list(extra.get("paragraph", []))
    char = [v for k, v in cfg["character"].items()] + list(extra.get("character", []))
    tpl = open(POSTIMPORT_TPL, encoding="utf-8").read()
    n_marks = len(ast_texts) - (1 if title else 0)
    ast = {"/*ASTCHAR*/": js(cfg["character"]["asterisk_ref"]), "/*ASTPARA*/": js(cfg["paragraph"]["asterisk_note"]),
           "/*ASTMARKS*/": str(n_marks), "/*ASTNOTES*/": str(len(ast_texts))}
    tpl = tpl.replace("/*TOTAL*/", str(total)).replace("/*PARA*/", js(para)).replace("/*CHAR*/", js(char))
    for k, v in ast.items():
        tpl = tpl.replace(k, v)
    open(os.path.join(outdir, stem + "_postimport.jsx"), "w", encoding="utf-8").write(tpl)
    if ast_texts:
        tpl = open(ASTERISK_TPL, encoding="utf-8").read()
        for k, v in ast.items():
            tpl = tpl.replace(k, v)
        tpl = tpl.replace("/*TITLE*/", "true" if title else "false").replace(
            "/*NOTES*/", ",\n".join("  " + js(t[:70]) for t in ast_texts))
        open(os.path.join(outdir, stem + "_gwiazdki.jsx"), "w", encoding="utf-8").write(tpl)


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md")
    ap.add_argument("--refs", action="append", default=None,
                    help="CSL-JSON; repeat to add srom-tlumacz's <id>_refs_tlum.json")
    ap.add_argument("--pair-src", help="frozen source (<id>_src.md): query rows that name a source note label "
                                        "get the printed note number")
    ap.add_argument("--out", default="build")
    ap.add_argument("--config", default=os.path.join(ROOT, "config", "styles.json"))
    ap.add_argument("--draft", action="store_true", help="allow [BRAK …] placeholders")
    ap.add_argument("--proof", action="store_true",
                    help="reading proof: always writes <stem>_korekta.docx with the full apparatus rendered, even with "
                         "errors (e.g. an English source before translation); not for InDesign")
    ap.add_argument("--queries", action="append", default=[],
                    help="extra query rows (CSV adresat;rodzaj;przypis;dzieło;szczegóły, e.g. from srom-tlumacz) merged into _pytania")
    a = ap.parse_args()

    os.makedirs(a.out, exist_ok=True)
    stem = os.path.splitext(os.path.basename(a.md))[0]
    work = tempfile.mkdtemp()
    cfg = json.load(open(a.config, encoding="utf-8"))
    ITALIC_STYLE[0] = cfg["character"]["italic"]
    BIB_NUMBERS[0] = bool(cfg.get("bibliography", {}).get("section_numbers", False))
    report = {"errors": [], "warnings": [], "flags": []}

    refs_list = []
    for rp in a.refs or []:
        refs_list += json.load(open(rp, encoding="utf-8"))
    refs_path = os.path.join(work, "refs.json")
    json.dump(refs_list, open(refs_path, "w", encoding="utf-8"), ensure_ascii=False)
    refs = {r["id"]: r for r in refs_list}
    dup = [k for k, v in Counter(r["id"] for r in refs_list).items() if v > 1]
    if dup:
        report["errors"].append(f"duplicate ids in refs: {dup}")

    md_text = open(a.md, encoding="utf-8").read()
    sys.path.insert(0, os.path.join(ROOT, "scripts"))
    import check as integrity
    md_text, comments = integrity.strip_comments(md_text)
    if comments:
        report["warnings"].append(f"{len(comments)} comment(s) removed from the text (never printed, never blocking): "
                                  + " | ".join(c[:60] for c in comments[:8]))

    # 0. integrity (check.py): note markers/definitions, keys, broken italics, author-date leftovers
    ierr, iwarn = [], []
    integrity.check_single(a.md, refs, ierr, iwarn)   # reads the file, strips comments itself
    report["errors"] += ["integrity: " + e for e in ierr]
    report["warnings"] += ["integrity: " + w for w in iwarn]

    # 1. citations in source order (on the comment-free text)
    clean_path = os.path.join(work, stem + ".clean.md")
    open(clean_path, "w", encoding="utf-8").write(md_text)
    pre, _ = pandoc_json(clean_path)
    cited = []
    def grab(x):
        if x.get("t") == "Cite":
            for c in x["c"][0]:
                if c["citationId"] not in cited:
                    cited.append(c["citationId"])
    cite_count = Counter()
    def count(x):
        if x.get("t") == "Cite":
            for c in x["c"][0]:
                cite_count[c["citationId"]] += 1
    walk(pre["blocks"], grab)
    walk(pre["blocks"], count)
    missing = [k for k in cited if k not in refs]
    if missing:
        report["errors"].append(f"citation keys not in refs: {missing}")
    uncited = [k for k in refs if k not in cited]
    if uncited:
        report["warnings"].append(f"refs never cited (not printed in bibliography): {uncited}")
    cited = [k for k in cited if k in refs]

    # 2. bibliography sections
    composed = build_bibliography(md_text, refs, cited, refs_path, work, report)
    comp_path = os.path.join(work, stem + ".md")
    open(comp_path, "w", encoding="utf-8").write(composed)

    # 3. render twice through the Lua filter: normal CSL, and CSL with Ibidem disabled (fallback forms)
    refdoc = os.path.join(work, "reference.docx")
    make_reference_docx(cfg, refdoc)
    docx_path = os.path.join(a.out, stem + ("_korekta.docx" if a.proof else ".docx"))
    if a.proof:
        a.draft = True
    meta = ["-M", "lang=pl-PL", "-M", "notes-after-punctuation=false", "-M", "suppress-bibliography=true",
            "-M", f"srom-config={os.path.abspath(a.config)}"]
    base = ["--citeproc", "--bibliography", refs_path, "--lua-filter", LUA, *meta]
    code, out_a, err = run(["pandoc", comp_path, "-f", FROM, *base, "--csl", CSL, "-t", "json"])
    nopage = []
    for ln in err.splitlines():
        if ln.startswith("SROM-WARN:"):
            report["warnings"].append(ln[10:].strip())
        elif ln.startswith("SROM-ERROR:"):
            report["errors"].append(ln[11:].strip())
        elif ln.startswith("SROM-NOPAGE:"):
            f = ln[12:].strip().split("\t")
            nopage.append((int(f[0]) if f[0].isdigit() else f[0], f[1], f[2], f[3] if len(f) > 3 else ""))
        elif ln.strip() and "SROM:" not in ln:
            report["errors"].append("pandoc: " + ln.strip())
    notes, rows, forced, literal_ibid = [], [], [], []
    n_na, ast_texts = 0, []
    n_title = 1 if re.search(r"^:::\s*\{?\.?przypis-tytulowy", md_text, re.M) else 0
    ok = False
    if code:
        report["errors"].append("pandoc/Lua filter failed — DOCX not produced")
    else:
        A = json.loads(out_a)
        B, _ = pandoc_json(comp_path, base + ["--csl", noibid_csl(work)])
        na, nb = notes_text(A), notes_text(B)
        pre_all = pre_notes(pandoc_json(comp_path)[0])
        # §7.1 non-author notes (– przyp. tłum./red.) leave the numbered sequence (Lua pass F); remember which
        # numbered note follows one: citeproc computed its Ibidem against the non-author note
        pre, after_na, prev_na = [], [], False
        for p_ in pre_all:
            if integrity.na_kind(p_):
                prev_na = True
                continue
            pre.append(p_)
            after_na.append(prev_na)
            prev_na = False
        n_na = len(pre_all) - len(pre)
        if len(pre) != len(na):
            report["errors"].append(f"internal: note sequence mismatch ({len(pre)} before / {len(na)} after citeproc)")
        # §7.3 Ibidem: kept only where it is unambiguous and grammatical; else the short form is printed
        for i, (x, y) in enumerate(zip(na, nb), 1):
            if "*Ibidem*" not in x or x == y:
                continue
            mid = not re.match(r"^(?:(?:Zob|Por)\. (?:też )?|Np\. )?\*Ibidem\*", x) and \
                not re.search(r"(?:; |[Zz]ob\. (?:też )?|[Pp]or\. (?:też )?|np\. )\*Ibidem\*", x)
            prev_extra = i >= 2 and i - 2 < len(pre) and extra_source(pre[i - 2])
            cur_extra = i - 1 < len(pre) and extra_source(pre[i - 1])
            after = i - 1 < len(after_na) and after_na[i - 1]
            if mid or prev_extra or cur_extra or after:
                why = ("Ibidem inside a sentence" if mid else "previous note also cites another (literal) source" if prev_extra
                       else "this note also cites another (literal) source before it" if cur_extra
                       else "the note before it in the text is a non-author (asterisk) note")
                forced.append((i, x, y, why))
            else:
                rows.append((i, x, y))
        literal_ibid = [i for i, x in enumerate(na, 1) if re.search(r"Ibidem|Ibid\.", x) and all(i != r[0] for r in rows + forced)]
        if forced:
            nda, ndb = note_dicts(A), note_dicts(B)
            for i, _, _, _ in forced:
                nda[i - 1]["c"] = ndb[i - 1]["c"]
        # §7.3: no Ibidem inside a non-author note (it is printed apart from the numbered series)
        ast_style = cfg["paragraph"]["asterisk_note"]
        def ast_divs(doc):
            return [b for b in doc["blocks"] if b["t"] == "Div" and dict(b["c"][0][2]).get("custom-style") == ast_style]
        for x_, y_ in zip(ast_divs(A), ast_divs(B)):
            if "Ibidem" in ser(x_["c"][1]) and x_ != y_:
                x_["c"] = copy.deepcopy(y_["c"])
                report["warnings"].append("Ibidem in a non-author note replaced by the short form: " + ser(y_["c"][1])[:90])
        ast_texts = [ser(b["c"][1]) for b in ast_divs(A) if ser(b["c"][1]).startswith("* ")]
        merged = os.path.join(work, stem + ".json")
        open(merged, "w", encoding="utf-8").write(json.dumps(A))
        notes = notes_text(A)
        code_d, _, err_d = run(["pandoc", "-f", "json", merged, "--reference-doc", refdoc, "-o", docx_path])
        if code_d:
            report["errors"].append("pandoc DOCX writer failed: " + err_d.strip()[:300])
        _, plain, _ = run(["pandoc", "-f", "json", merged, "-t", "plain", "--wrap=none"])
        open(os.path.join(a.out, stem + ".txt"), "w", encoding="utf-8").write(plain)
        brak = sorted(set(re.findall(r"\[BRAK [A-ZĄĆĘŁŃÓŚŹŻ ]+\]", plain)))
        if brak:
            brak_notes = [i for i, n in enumerate(notes, 1) if "[BRAK" in n]
            msg = f"missing data placeholders {brak} in notes {brak_notes} — ask the author (kanon §0; see {stem}_pytania.md)"
            (report["warnings"] if a.draft else report["errors"]).append(msg)
        if os.path.exists(docx_path) and not code_d:
            id2name = postprocess_docx(docx_path, cfg, report)
            ok = verify_docx(docx_path, cfg, id2name, len(notes), report, (n_na, n_title))
            for n_, ok_, d_ in report.get("verify", []):
                if not ok_:
                    report["errors"].append(f"DOCX verification failed — {n_}: {d_}")

    # 5. InDesign scripts + author query sheet
    if not a.proof:
        write_jsx(rows, len(notes), cfg, a.out, stem, ast_texts, bool(n_title))
    label_map = {}
    if a.pair_src:
        try:
            label_map = integrity.printed_numbers(a.pair_src, a.md)
        except Exception as e:
            report["warnings"].append(f"--pair-src: note alignment failed ({e}); query rows keep source labels")
    write_queries(a.out, stem, nopage, notes, refs, cited, cite_count, refs_path, work, report, a.queries, label_map)
    if n_na or n_title:
        report["warnings"].append(
            f"asterisk series (kanon § 7.1): {'title note + ' if n_title else ''}{n_na} translator/editorial note(s) are "
            f"paragraphs in '{cfg['paragraph']['asterisk_note']}' at the end of the DOCX, marked * in the text — set them "
            f"by hand above the numbered notes; after layout run {stem}_gwiazdki.jsx for the asterisks per page")

    # 6. lint (the srom-kanon skill's linter — required)
    import kanon_path
    lint_out = ""
    kdir = kanon_path.find_kanon()
    report["linter"] = None
    if not kdir:
        report["errors"].append(kanon_path.MISSING)
    else:
        report["linter"] = f"{kanon_path.linter(kdir)} (Kanon v{kanon_path.kanon_version(kdir) or '?'})"
        _, lint_out, lint_err = run([sys.executable, kanon_path.linter(kdir), os.path.join(a.out, stem + ".txt")])
        if not lint_out:
            report["errors"].append("kanon linter produced no output — cannot certify: " + lint_err.strip()[:300])
    if "--- ERROR ---" in lint_out:
        n_err = len(re.findall(r"^\d+:\d+\s+\[", lint_out.split("--- ERROR ---")[1].split("--- WARN ---")[0], re.M))
        report["errors"].append(f"kanon linter: {n_err} ERROR(s) in rendered text — see Lint section")

    success = ok and not report["errors"]
    if a.proof:
        rep_head = "PROOF — reading copy with the full apparatus; not for InDesign"
    rep = [f"# Build report — {stem}", "",
           f"**RESULT: {rep_head if a.proof else ('PASS — ready for InDesign' if success else 'FAIL — do not typeset')}**", "",
           f"- footnotes: {len(notes)} · citation keys used: {len(cited)} · bibliography sections: {report.get('bib_sections')}",
           f"- pandoc {PANDOC_VERSION} · CSL {os.path.basename(CSL)} · config {os.path.basename(a.config)} · linter {report.get('linter') or 'none'}",
           f"- Ibidem notes to check after layout: {len(rows)} (run {stem}_ibidem.jsx)"
           + (f"; literal (non-CSL) Ibidem notes, check by hand: {literal_ibid}" if literal_ibid else ""),
           "", "## Errors"] + ([f"- {e}" for e in report["errors"]] or ["- none"])
    rep += ["", "## Warnings (review)"] + ([f"- {w}" for w in report["warnings"]] or ["- none"])
    rep += ["", "## DOCX verification"] + [f"- [{'x' if ok_ else ' '}] {n}: {d}" for n, ok_, d in report.get("verify", [])]
    rep += ["", "## Ibidem replaced by the short form at build time (§7.3)"]
    rep += [f"- {i}: {y}  — {why}" for i, x, y, why in forced] or ["- none"]
    rep += ["", "## Ibidem map (note → form to use if it lands on a different column)"]
    rep += [f"- {i}: {x}  ⇒  {y}" for i, x, y in rows] or ["- none"]
    rep += ["", "## Lint (srom-kanon lint_srom.py on rendered text)", "```", lint_out.strip()[:6000], "```"]
    open(os.path.join(a.out, stem + "_report.md"), "w", encoding="utf-8").write("\n".join(rep) + "\n")
    shutil.rmtree(work)
    if a.proof:
        print(f"PROOF — {docx_path} ({len(report['errors'])} issue(s) listed in {stem}_report.md)")
        sys.exit(0 if os.path.exists(docx_path) else 1)
    print(("PASS" if success else "FAIL") + f" — {os.path.join(a.out, stem + '_report.md')}")
    sys.exit(0 if success else 1)


if __name__ == "__main__":
    main()
