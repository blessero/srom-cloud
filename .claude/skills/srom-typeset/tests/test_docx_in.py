"""G13: DOCX import — footnotes kept and moved under their paragraph, italics and small caps kept,
tracked changes/comments counted, hand-typed superscript numbers, manual headings and numbered
lists flagged; the chain import -> normalize -> build stops on the forbidden 'Tamże'."""
import os, re, sys, json, subprocess, tempfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
S = os.path.join(ROOT, "scripts")
d = tempfile.mkdtemp()
res = []
def t(name, ok, detail=""):
    res.append(ok); print(("PASS " if ok else "FAIL ") + name + ("" if ok else "\n   " + str(detail)[:1500]))
docx = os.path.join(d, "in.docx")
subprocess.run(["pandoc", os.path.join(ROOT, "tests", "fixtures", "docx_in_src.md"), "-o", docx], check=True)
r = subprocess.run([sys.executable, os.path.join(S, "docx_in.py"), docx, "-o", os.path.join(d, "a.md")], capture_output=True, text=True)
md = open(os.path.join(d, "a.md"), encoding="utf-8").read()
rep = open(os.path.join(d, "a_import.md"), encoding="utf-8").read()
blocks = [b for b in re.split(r"\n\s*\n", md) if b.strip()]
t("footnote definitions directly under their paragraph (incl. 2nd paragraph of note 2)",
  blocks[1].startswith("Pierwszy akapit") and blocks[2].startswith("[^1]:") and blocks[3].startswith("[^2]:") and blocks[4].startswith("    Drugi akapit"), blocks[:6])
t("italics and small caps preserved", "*kursywą*" in md and "[Ficowski]{.smallcaps}" in md and "*Cyganie na polskich drogach*" in md, md)
t("tracked change accepted and counted, comment dropped and counted", "wstawką" in md and "1 insertions" in rep and "comments dropped: 1" in rep, rep)
t("hand-typed superscript digit flagged as possible fake footnote", "FAKE-NOTE" in rep, rep)
t("bold-only paragraph flagged as manual heading", "BOLDPARA" in rep, rep)
t("numbered list kept as a list (kanon allows numbered lists)", "1.  pierwszy punkt" in md or "1. pierwszy punkt" in md, md)
t("link reduced to text, URL listed", "oraz link." in md and "https://example.org" in rep, md)
n_md = os.path.join(d, "n.md")
subprocess.run([sys.executable, os.path.join(S, "normalize.py"), os.path.join(d, "a.md"), "-o", n_md, "--log", os.path.join(d, "n.log")], capture_output=True)
b = subprocess.run([sys.executable, os.path.join(S, "build.py"), n_md, "--out", os.path.join(d, "b")], capture_output=True, text=True)
brep = open(os.path.join(d, "b", "n_report.md"), encoding="utf-8").read()
t("import -> normalize -> build: unkeyed 'Tamże' blocks the build (kanon §7.3)", b.returncode == 1 and "kanon linter:" in brep and "ABBR" in brep, brep[:1500])
# ---- Zotero citation fields are harvested; the author's reference list is split off to _bib.txt
import docx as _docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
D = _docx.Document()
D.add_heading("1. Wstęp", level=1)
p = D.add_paragraph("Tekst przed cytowaniem ")
cit = {"citationID": "abc", "properties": {"formattedCitation": "(Ficowski 1985, 15)", "plainCitation": "(Ficowski 1985, 15)", "noteIndex": 0},
       "citationItems": [{"id": 123, "uris": ["http://zotero.org/users/1/items/ABCD1234"], "locator": "15", "label": "page",
                          "itemData": {"id": 123, "type": "book", "title": "Cyganie na polskich drogach", "publisher": "Wydawnictwo Literackie",
                                       "publisher-place": "Kraków", "author": [{"family": "Ficowski", "given": "Jerzy"}], "issued": {"date-parts": [["1985"]]},
                                       "ISBN": "83-08-01234-5"}}],
       "schema": "https://github.com/citation-style-language/schema/raw/master/csl-citation.json"}
def fld(par, instr, shown):
    for kind, txt in (("begin", None), ("instr", instr), ("separate", None), ("text", shown), ("end", None)):
        r = par.add_run(txt if kind == "text" else None)
        if kind == "instr":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
        elif kind != "text":
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r._r.append(e)
fld(p, " ADDIN ZOTERO_ITEM CSL_CITATION " + json.dumps(cit) + " ", "(Ficowski 1985, 15)")
p.add_run(" i dalej.")
D.add_heading("Bibliografia", level=1)
D.add_paragraph("Ficowski J., Cyganie na polskich drogach, Kraków 1985.")
D.add_paragraph("Mróz L., Dzieje Cyganów-Romów w Rzeczypospolitej, Warszawa 2001.")
zd = os.path.join(d, "z.docx"); D.save(zd)
r = subprocess.run([sys.executable, os.path.join(S, "docx_in.py"), zd, "-o", os.path.join(d, "z.md")], capture_output=True, text=True)
zrep = open(os.path.join(d, "z_import.md"), encoding="utf-8").read()
cited = open(os.path.join(d, "z_cited.md"), encoding="utf-8").read() if os.path.exists(os.path.join(d, "z_cited.md")) else ""
zrefs = json.load(open(os.path.join(d, "z_refs.json"), encoding="utf-8")) if os.path.exists(os.path.join(d, "z_refs.json")) else []
t("Zotero field harvested: citation keyed surnameYEAR with Polish locator", "[@ficowski1985, s. 15]" in cited and [x["id"] for x in zrefs] == ["ficowski1985"], zrep + cited)
t("author's reference list removed from the text and written to _bib.txt", "Bibliografia" not in cited and open(os.path.join(d, "z_bib.txt"), encoding="utf-8").read().count("\n") == 2, cited)
b = subprocess.run([sys.executable, os.path.join(S, "build.py"), os.path.join(d, "z_cited.md"), "--refs", os.path.join(d, "z_refs.json"), "--out", os.path.join(d, "zb")], capture_output=True, text=True)
zbt = open(os.path.join(d, "zb", "z_cited.txt"), encoding="utf-8").read() if os.path.exists(os.path.join(d, "zb", "z_cited.txt")) else ""
t("harvested article builds; the author-date citation becomes a SROM footnote", b.returncode == 0 and "cytowaniem[1] i dalej" in zbt and "J. Ficowski, Cyganie na polskich drogach, Wydawnictwo Literackie, Kraków 1985, s. 15." in zbt, b.stdout + zbt)

D2 = _docx.Document(); D2.add_paragraph("(1)\tMe\tdikhav"); D2.add_paragraph("Zwykły akapit."); td = os.path.join(d, "tabs.docx"); D2.save(td)
r = subprocess.run([sys.executable, os.path.join(S, "docx_in.py"), td, "-o", os.path.join(d, "tabs.md")], capture_output=True, text=True)
trep = open(os.path.join(d, "tabs_import.md"), encoding="utf-8").read()
t("tab-aligned paragraph (interlinear example typed in Word) flagged", "TABS" in trep and "(1)" in trep and trep.count("TABS") == 1 and "IMPORT CHECK" in r.stdout, trep)

n, ok = len(res), sum(res)
print(f"DOCXIN ALL PASS {n}/{n}" if ok == n else f"DOCXIN FAILED {n - ok}/{n}")
